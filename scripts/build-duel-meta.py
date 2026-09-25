#!/usr/bin/env python3
"""
Génère src/data/duel-meta.json à partir du classeur d'analyse
analysis/duelcommander/duelcommander_cartes_incontournables.xlsx
(25/09/2026, constructeur compétitif — voir competitive-builder.ts).

Ce classeur agrège 82 decklists Duel Commander de tournois réels publiées
sur mtgtop8.com (01/09/2026 → 04/09/2026, 20 tournois). On n'en garde que
ce qui est fiable pour notre usage :
- le NOM de chaque carte et sa fréquence de présence (transcrits depuis
  les decklists — fiables) ;
- les commandants joués (idem).
On IGNORE volontairement les colonnes couleur/rareté du classeur : elles
ont été estimées sans accès à Scryfall (voir la feuille « Résumé ») ;
l'identité couleur et la légalité sont revérifiées à l'exécution via
Scryfall par le constructeur.

Seules les cartes présentes dans au moins 2 decks sont gardées (une
présence dans 1 seul deck sur 82 n'est pas un signal de méta).

Usage : python3 scripts/build-duel-meta.py  (nécessite openpyxl)
"""
import json
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "analysis/duelcommander/duelcommander_cartes_incontournables.xlsx"
OUT = ROOT / "src/data/duel-meta.json"

wb = openpyxl.load_workbook(SRC, data_only=True)

summary = {r[1]: r[2] for r in wb["Résumé"].iter_rows(values_only=True) if r and r[1]}
deck_count = int(summary.get("Decks analysés", 0))

cards = {}
for row in wb["Toutes les cartes"].iter_rows(min_row=2, values_only=True):
    name, freq, share = row[0], row[1], row[2]
    if not name or not freq or freq < 2:
        continue
    cards[str(name).strip()] = round(float(share), 3)

commanders = {}
for row in wb["Decks (sources)"].iter_rows(min_row=2, values_only=True):
    if not row or not row[1]:
        continue
    # Partenaires notés "A / B" : chaque nom compte (un partenaire peut
    # aussi être commandant seul — le constructeur ne gère qu'un commandant).
    for name in str(row[1]).split(" / "):
        name = name.strip()
        commanders[name] = commanders.get(name, 0) + 1

data = {
    "source": "mtgtop8.com — 82 decks Duel Commander de tournoi, analyse analysis/duelcommander/duelcommander_cartes_incontournables.xlsx",
    "period": summary.get("Période couverte"),
    "deckCount": deck_count,
    "cards": dict(sorted(cards.items(), key=lambda kv: -kv[1])),
    "commanders": dict(sorted(commanders.items(), key=lambda kv: -kv[1])),
}
OUT.write_text(json.dumps(data, ensure_ascii=False, indent=0) + "\n", encoding="utf-8")
print(f"{len(cards)} cartes, {len(commanders)} commandants -> {OUT}")
